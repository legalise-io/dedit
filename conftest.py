from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from backend import main


@pytest.fixture
def client():
    main.documents.clear()
    main.templates.clear()
    with TestClient(main.app, raise_server_exceptions=False) as instance:
        yield instance
    main.documents.clear()
    main.templates.clear()


@pytest.fixture
def docx_bytes():
    doc = Document()
    doc.add_paragraph("A café & a contract <draft>.")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


@pytest.fixture
def tracked_tabs_docx():
    """Portable replacement for the ignored GliderContract.docx fixture."""
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("plain\ttab")
    for kind, text in [("ins", "new"), ("del", "old")]:
        revision = OxmlElement(f"w:{kind}")
        for attr, value in {"id": kind, "author": "Alice", "date": "2025-01-01T00:00:00Z"}.items():
            revision.set(qn(f"w:{attr}"), value)
        run = OxmlElement("w:r")
        prop = OxmlElement("w:rPr")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "dotted")
        prop.append(underline)
        run.append(prop)
        txt = OxmlElement("w:delText" if kind == "del" else "w:t")
        txt.text = text
        run.append(txt)
        run.append(OxmlElement("w:tab"))
        revision.append(run)
        para._p.append(revision)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
