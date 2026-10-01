import copy
import json
from io import BytesIO

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches
from docx2tiptap import comments_to_dict, create_docx_from_tiptap, elements_to_dict, parse_docx, to_tiptap
from docx2tiptap.comments_parser import extract_comments_from_docx
from docx2tiptap.docx_exporter import DocxExporter
from docx2tiptap.models import Paragraph, TextRun, Section
from docx2tiptap.utils import base64_to_element, element_to_base64

from .helpers import NS, document, imported, paragraph, revision, save, text, walk, xml


@pytest.mark.parametrize("value", ["Hello", " café & <draft> ", "中文 😀", "line\tend"])
def test_text_roundtrip(value):
    source = document(paragraph(text(value)))
    result = imported(create_docx_from_tiptap(source).getvalue())
    assert "".join(n.get("text", "") for n in walk(result)) == value


@pytest.mark.parametrize("marks", [[{"type": "bold"}], [{"type": "italic"}], [{"type": "bold"}, {"type": "italic"}]])
def test_basic_marks_roundtrip(marks):
    result = imported(create_docx_from_tiptap(document(paragraph(text("word", *marks)))).getvalue())
    actual = next(n for n in walk(result) if n.get("text") == "word")
    assert {m["type"] for m in marks} <= {m["type"] for m in actual["marks"]}


@pytest.mark.parametrize("level", range(1, 7))
def test_headings_roundtrip(level):
    source = document({"type": "heading", "attrs": {"level": level}, "content": [text("Title")]})
    result = imported(create_docx_from_tiptap(source).getvalue())
    assert result["content"][0]["type"] == "heading"
    assert result["content"][0]["attrs"]["level"] == level


@pytest.mark.parametrize("kind", ["insertion", "deletion"])
def test_revision_metadata_and_formatting_roundtrip(kind):
    source = document(paragraph(text("tracked", revision(kind), {"type": "bold"})))
    result = imported(create_docx_from_tiptap(source).getvalue())
    node = next(n for n in walk(result) if n.get("text") == "tracked")
    mark = next(m for m in node["marks"] if m["type"] == kind)
    assert mark["attrs"]["author"] == "Alice"
    assert mark["attrs"]["date"] == "2025-01-01T00:00:00Z"
    assert any(m["type"] == "bold" for m in node["marks"])


@pytest.mark.parametrize("kind", [None, "page", "column"])
def test_break_roundtrip(kind):
    source = document(paragraph(text("a"), {"type": "hardBreak", "attrs": {"breakType": kind}}, text("b")))
    result = imported(create_docx_from_tiptap(source).getvalue())
    br = next(n for n in walk(result) if n["type"] == "hardBreak")
    assert br.get("attrs", {}).get("breakType") == kind


def test_empty_document_is_editable():
    result = to_tiptap([])
    assert result == document({"type": "paragraph"})
    assert Document(create_docx_from_tiptap(result)) is not None


def test_import_does_not_mutate_intermediate_models():
    elements = [Paragraph(runs=[TextRun(text="hello", bold=True)])]
    original = copy.deepcopy(elements)
    result = to_tiptap(elements)
    assert elements == original
    json.dumps(elements_to_dict(elements))
    json.dumps(result)


def test_export_does_not_mutate_input():
    source = document(paragraph(text("body")), {"type": "rawStylesStorage", "attrs": {"data": "{}"}})
    before = copy.deepcopy(source)
    create_docx_from_tiptap(source)
    assert source == before


def test_reusing_exporter_resets_revision_counter():
    exporter = DocxExporter()
    source = document(paragraph(text("new", revision("insertion"))))
    ids = [xml(exporter.export(source)).xpath("//w:ins/@w:id", namespaces=NS) for _ in range(2)]
    assert ids == [["1"], ["1"]]


def test_template_preserves_header_footer_and_page_layout():
    template = Document()
    template.add_paragraph("old content")
    template.sections[0].header.paragraphs[0].text = "Header"
    template.sections[0].footer.paragraphs[0].text = "Footer"
    template.sections[0].left_margin = Inches(1.75)
    output = Document(create_docx_from_tiptap(document(paragraph(text("new content"))), template_bytes=save(template)))
    assert [p.text for p in output.paragraphs] == ["new content"]
    assert output.sections[0].header.paragraphs[0].text == "Header"
    assert output.sections[0].footer.paragraphs[0].text == "Footer"
    assert output.sections[0].left_margin == Inches(1.75)


def test_comments_on_plain_text_roundtrip():
    mark = {"type": "comment", "attrs": {"commentId": "c1"}}
    data = create_docx_from_tiptap(document(paragraph(text("anchor", mark))), [{"id": "c1", "author": "Bob", "text": "Check this", "initials": "B"}]).getvalue()
    elements, comments, numbering = parse_docx(data)
    serialized = comments_to_dict(comments)
    assert len(serialized) == 1
    assert serialized[0]["text"] == "Check this"
    assert serialized[0]["author"] == "Bob"
    assert serialized[0]["initials"] == "B"
    assert any(m["type"] == "comment" for n in walk(to_tiptap(elements, comments, numbering)) for m in n.get("marks", []))


def test_missing_comment_data_does_not_create_phantom_comment():
    mark = {"type": "comment", "attrs": {"commentId": "absent"}}
    output = create_docx_from_tiptap(document(paragraph(text("anchor", mark)))).getvalue()
    assert extract_comments_from_docx(output) == {}


@pytest.mark.parametrize("payload", [b"", b"not a zip"])
def test_invalid_docx_rejected(payload):
    with pytest.raises(Exception):
        parse_docx(payload)


def test_invalid_comment_container_returns_empty():
    assert extract_comments_from_docx(b"bad zip") == {}


def test_table_content_and_horizontal_merge_roundtrip():
    source = Document()
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
    table.cell(1, 0).text = "Left"
    table.cell(1, 1).text = "Right"
    raw = save(source)
    result = imported(create_docx_from_tiptap(imported(raw), template_bytes=raw).getvalue())
    table_node = next(n for n in walk(result) if n["type"] == "table")
    assert table_node["content"][0]["content"][0]["attrs"]["colspan"] == 2
    assert {n.get("text") for n in walk(table_node) if n["type"] == "text"} == {"Merged", "Left", "Right"}


def test_vertical_merge_roundtrip():
    source = Document()
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Vertical"
    table.cell(0, 1).text = "Top"
    table.cell(1, 1).text = "Bottom"
    raw = save(source)
    result = imported(create_docx_from_tiptap(imported(raw), template_bytes=raw).getvalue())
    table_node = next(n for n in walk(result) if n["type"] == "table")
    assert table_node["content"][0]["content"][0]["attrs"]["rowspan"] == 2
    assert "Bottom" in {n.get("text") for n in walk(table_node)}


def test_nested_tables_roundtrip():
    source = Document()
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Outer"
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Inner"
    raw = save(source)
    result = imported(create_docx_from_tiptap(imported(raw), template_bytes=raw).getvalue())
    assert len([n for n in walk(result) if n["type"] == "table"]) == 2
    assert {"Outer", "Inner"} <= {n.get("text") for n in walk(result)}


def test_direct_paragraph_formatting_survives_unedited_roundtrip():
    source = Document()
    para = source.add_paragraph("Indented")
    para.paragraph_format.left_indent = Inches(0.5)
    raw = save(source)
    doc = imported(raw)
    storage = next(n for n in doc["content"] if n["type"] == "rawStylesStorage")
    assert f"para:{doc['content'][0]['attrs']['id']}:pPr" in json.loads(storage["attrs"]["data"])
    result = Document(create_docx_from_tiptap(doc, template_bytes=raw))
    assert result.paragraphs[0].paragraph_format.left_indent == Inches(0.5)


def test_numbered_paragraphs_preserve_definition_reference():
    source = Document()
    source.add_paragraph("First", style="List Number")
    source.add_paragraph("Second", style="List Number")
    raw = save(source)
    doc = imported(raw)
    paras = [n for n in doc["content"] if n["type"] == "paragraph"]
    assert all(n["attrs"].get("numId") for n in paras)
    assert [n["attrs"]["styleNumbering"] for n in paras] == ["1.", "2."]
    output = Document(create_docx_from_tiptap(doc, template_bytes=raw))
    assert [p.text for p in output.paragraphs] == ["First", "Second"]


def test_section_nodes_export_child_content():
    source = document({"type": "section", "content": [paragraph(text("inside"))]})
    assert Document(create_docx_from_tiptap(source)).paragraphs[0].text == "inside"


def test_raw_xml_helpers_roundtrip():
    element = Document().add_paragraph("x")._p
    restored = base64_to_element(element_to_base64(element))
    assert restored.tag == element.tag
    assert restored.xpath(".//w:t", namespaces=NS)[0].text == "x"


@pytest.mark.parametrize("payload", ["!invalid!", "", None])
def test_bad_raw_xml_returns_none(payload):
    assert base64_to_element(payload) is None


def test_rich_table_attributes_export_and_import():
    border = {'style': 'double', 'width': 8, 'color': '112233'}
    attrs = {'backgroundColor': '#abcdef', 'verticalAlign': 'center', 'colwidth': [1800],
             'textAlign': 'right', 'borders': {side: border for side in ['top', 'bottom', 'left', 'right']}}
    table = {'type': 'table', 'attrs': {'alignment': 'right', 'colwidths': [1800]}, 'content': [
        {'type': 'tableRow', 'content': [{'type': 'tableCell', 'attrs': attrs, 'content': [paragraph(text('styled'))]}]}]}
    output = create_docx_from_tiptap(document(table)).getvalue()
    tree = xml(output)
    assert tree.xpath('//w:tcPr/w:shd/@w:fill', namespaces=NS) == ['abcdef']
    assert tree.xpath('//w:tcPr/w:tcBorders/w:right/@w:sz', namespaces=NS) == ['8']
    parsed = imported(output)
    actual = next(n for n in walk(parsed) if n['type'] in ('tableCell', 'tableHeader'))['attrs']
    assert actual['backgroundColor'] == '#abcdef'
    assert actual['verticalAlign'] == 'center'
    assert actual['textAlign'] == 'right'
    assert actual['borders']['right'] == border
    rebuilt = xml(create_docx_from_tiptap(parsed))
    assert rebuilt.xpath('//w:tcPr/w:shd/@w:fill', namespaces=NS) == ['abcdef']


@pytest.mark.parametrize('change', [
    {'id': '42', 'author': 'Bob', 'date': '2025-01-01T00:00:00Z', 'oldStyle': 'ListNumber', 'oldNumIlvl': 0},
    {'oldNumIlvl': 0},
])
def test_numbered_paragraph_format_change_roundtrip(change):
    source = document(paragraph(text('changed'), numId='7', numIlvl=1, formatChange=change))
    output = create_docx_from_tiptap(source).getvalue()
    tree = xml(output)
    assert tree.xpath('//w:pPrChange/w:pPr/w:numPr/w:ilvl/@w:val', namespaces=NS) == ['0']
    parsed = imported(output)['content'][0]['attrs']['formatChange']
    assert parsed['oldNumIlvl'] == 0
    assert parsed['author'] == change.get('author', 'Unknown')
