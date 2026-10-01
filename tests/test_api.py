from io import BytesIO

import pytest
from docx import Document

from .helpers import document, paragraph, text


def upload(client, data, name="contract.docx", path="/upload"):
    return client.post(path, files={"file": (name, data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})


def test_health_and_empty_lists(client):
    assert client.get("/").json()["status"] == "ok"
    assert client.get("/documents").json() == []
    assert client.get("/templates").json() == []


def test_upload_converts_and_lists_document(client, docx_bytes):
    response = upload(client, docx_bytes)
    assert response.status_code == 200
    doc = response.json()
    assert doc["tiptap"]["type"] == "doc"
    assert doc["comments"] == []
    assert doc["intermediate"]
    assert client.get("/documents").json() == [{"id": doc["id"], "filename": "contract.docx"}]


@pytest.mark.parametrize("name", ["bad.txt", "", "bad.pdf"])
def test_upload_filename_validation(client, name):
    assert upload(client, b"bad", name).status_code in (400, 422)


def test_corrupt_upload_not_stored(client):
    assert upload(client, b"bad").status_code == 500
    assert client.get("/documents").json() == []


def test_missing_document(client):
    assert client.get("/documents/missing").status_code == 404


@pytest.mark.parametrize("filename", ["report", "report.pdf", "report.docx"])
def test_export_returns_valid_docx_and_filename(client, filename):
    response = client.post("/export", json={"tiptap": document(paragraph(text("Exported"))), "filename": filename})
    assert response.status_code == 200
    assert "wordprocessingml.document" in response.headers["content-type"]
    assert 'filename="report.docx"' in response.headers["content-disposition"]
    assert Document(BytesIO(response.content)).paragraphs[0].text == "Exported"


@pytest.mark.parametrize("template,field", [("original", "document_id"), ("custom", "template_id")])
def test_export_requires_template_identifier(client, template, field):
    payload = {"tiptap": document(paragraph(text("x"))), "template": template}
    assert client.post("/export", json=payload).status_code == 400
    payload[field] = "missing"
    assert client.post("/export", json=payload).status_code == 404


@pytest.mark.parametrize("payload", [{}, {"tiptap": []}, {"tiptap": {}, "template": "unknown"}])
def test_export_request_validation(client, payload):
    assert client.post("/export", json=payload).status_code == 422


def test_export_uses_original_template(client, docx_bytes):
    doc_id = upload(client, docx_bytes).json()["id"]
    response = client.post("/export", json={"tiptap": document(paragraph(text("Changed"))), "template": "original", "document_id": doc_id})
    assert response.status_code == 200
    assert Document(BytesIO(response.content)).paragraphs[0].text == "Changed"


def test_template_upload_export_and_delete(client, docx_bytes):
    response = upload(client, docx_bytes, path="/templates/upload")
    assert response.status_code == 200
    template_id = response.json()["id"]
    assert client.get("/templates").json() == [{"id": template_id}]
    response = client.post("/export", json={"tiptap": document(paragraph(text("Custom"))), "template": "custom", "template_id": template_id})
    assert response.status_code == 200
    assert Document(BytesIO(response.content)).paragraphs[0].text == "Custom"
    assert client.delete(f"/templates/{template_id}").status_code == 200
    assert client.delete(f"/templates/{template_id}").status_code == 404
    assert client.get("/templates").json() == []


def test_template_filename_validation(client):
    assert upload(client, b"x", "bad.txt", "/templates/upload").status_code == 400


def test_cors_allows_demo_origin(client):
    response = client.options("/upload", headers={"Origin": "http://localhost:5173", "Access-Control-Request-Method": "POST"})
    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "http://localhost:5173"


def test_cors_rejects_other_origin(client):
    response = client.options("/upload", headers={"Origin": "https://example.com", "Access-Control-Request-Method": "POST"})
    assert response.status_code == 400
