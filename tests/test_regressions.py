"""Regression coverage for confirmed and repaired review defects."""
from io import BytesIO

import pytest
from docx import Document
from docx.shared import Inches
from docx2tiptap import create_docx_from_tiptap, parse_docx
from docx2tiptap.utils import element_to_base64

from .helpers import NS, document, imported, paragraph, revision, save, text, walk, xml


def test_uploaded_document_can_be_retrieved(client, docx_bytes):
    response = client.post("/upload", files={"file": ("test.docx", docx_bytes)})
    assert response.status_code == 200
    response = client.get(f"/documents/{response.json()['id']}")
    assert response.status_code == 200
    assert response.json()["tiptap"]["type"] == "doc"


def test_inserting_paragraph_does_not_move_original_indentation():
    original = Document()
    original.add_paragraph("Indented").paragraph_format.left_indent = Inches(0.5)
    data = save(original)
    doc = imported(data)
    doc["content"].insert(0, paragraph(text("New")))
    result = Document(create_docx_from_tiptap(doc, template_bytes=data))
    assert result.paragraphs[0].paragraph_format.left_indent is None
    assert result.paragraphs[1].paragraph_format.left_indent == Inches(0.5)


def test_comments_on_imported_formatted_text_survive_export():
    original = Document()
    run = original.add_paragraph().add_run("Bold anchor")
    run.bold = True
    original.add_comment(run, text="Comment", author="Alice")
    data = save(original)
    doc = imported(data)
    comments = [{"id": "0", "text": "Comment", "author": "Alice"}]
    output = create_docx_from_tiptap(doc, comments).getvalue()
    assert len(parse_docx(output)[1]) == 1


def test_removing_bold_in_editor_removes_bold_in_word():
    original = Document()
    original.add_paragraph().add_run("bold").bold = True
    doc = imported(save(original))
    for node in walk(doc):
        if node.get("type") == "text":
            node["marks"] = [m for m in node.get("marks", []) if m["type"] != "bold"]
    output = Document(create_docx_from_tiptap(doc))
    assert output.paragraphs[0].runs[0].bold is not True


def test_blank_paragraphs_are_preserved():
    original = Document()
    original.add_paragraph("Before")
    original.add_paragraph()
    original.add_paragraph("After")
    result = Document(create_docx_from_tiptap(imported(save(original))))
    assert [p.text for p in result.paragraphs] == ["Before", "", "After"]


def test_heading_revision_inside_table_is_preserved():
    heading = {"type": "heading", "attrs": {"level": 1}, "content": [text("Inserted", revision("insertion"))]}
    table = {"type": "table", "content": [{"type": "tableRow", "content": [{"type": "tableCell", "content": [heading]}]}]}
    output = xml(create_docx_from_tiptap(document(table)))
    assert output.xpath("//w:tbl//w:ins", namespaces=NS)


def test_repeated_text_has_comment_on_correct_occurrence():
    original = Document()
    para = original.add_paragraph()
    first = para.add_run("same")
    para.add_run("same")
    original.add_comment(first, text="Only first", author="Alice")
    doc = imported(save(original))
    text_nodes = [n for n in walk(doc) if n.get("type") == "text"]
    assert any(m["type"] == "comment" for m in text_nodes[0].get("marks", []))
    assert not any(m["type"] == "comment" for m in text_nodes[1].get("marks", []))


def test_explicitly_disabled_bold_is_not_imported_as_bold():
    original = Document()
    original.add_paragraph().add_run("plain").bold = False
    doc = imported(save(original))
    assert not any(m["type"] == "bold" for n in walk(doc) for m in n.get("marks", []))


def test_comments_on_tracked_text_survive_export():
    mark = {"type": "comment", "attrs": {"commentId": "c1"}}
    source = document(paragraph(text("new", revision("insertion"), mark)))
    output = create_docx_from_tiptap(source, [{"id": "c1", "text": "Check", "author": "Alice"}]).getvalue()
    assert len(parse_docx(output)[1]) == 1
